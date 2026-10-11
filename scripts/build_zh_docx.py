# -*- coding: utf-8 -*-
"""Build submission-ready Word file (manuscript_zh.docx) from assembled manuscript.md.

Cleans workflow artifacts (blockquote notes, reference verification annotations),
then renders A4 double-spaced manuscript with line numbers via python-docx.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

ROOT = Path(__file__).resolve().parents[1]
P = ROOT / "outputs" / "paper"
OUT = ROOT / "outputs" / "submission" / "manuscript_zh.docx"

SONG, HEI, ASCII = "宋体", "黑体", "Times New Roman"

# ---------- load & clean ----------
text = (P / "manuscript.md").read_text(encoding="utf-8")
lines = text.splitlines()

cleaned = []
for ln in lines:
    s = ln.rstrip()
    if s.strip().startswith(">"):
        continue  # workflow notes
    if s.strip() == "###":
        continue
    if s.startswith("## 参考文献骨架"):
        cleaned.append("## 参考文献")
        continue
    cleaned.append(s)
text = "\n".join(cleaned)

# strip verification/meta annotations inside fullwidth parens in reference lines
def strip_ref_notes(m: re.Match) -> str:
    inner = m.group(1)
    if re.search(r"核验|英文稿文献|会议论文集页码|活页指南", inner):
        return ""
    return m.group(0)

lines = []
in_refs = False
for ln in text.splitlines():
    if ln.startswith("## 参考文献"):
        in_refs = True
    elif ln.startswith("## "):
        in_refs = False
    if in_refs and ln.startswith(("1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.", "10.", "11.", "12.")):
        ln = re.sub(r"（([^）]*)）", strip_ref_notes, ln).rstrip()
    lines.append(ln)
text = "\n".join(lines)

# ---------- docx helpers ----------
doc = Document()

sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Mm(25.4)
sec.left_margin = sec.right_margin = Mm(31.7)
lnnum = sec._sectPr.makeelement(qn("w:lnNumType"), {qn("w:countBy"): "1", qn("w:restart"): "newPage"})
sec._sectPr.append(lnnum)

style = doc.styles["Normal"]
style.font.name = ASCII
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)


def set_run(run, size=12, bold=False, ea=SONG, ascii_=ASCII, sup=False):
    run.font.name = ascii_
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.superscript = sup
    run.font.color.rgb = None or run.font.color.rgb  # keep
    rPr = run._element.get_or_add_rPr()
    rF = rPr.get_or_add_rFonts()
    rF.set(qn("w:eastAsia"), ea)


TOKEN = re.compile(r"(<sup>.+?</sup>|\*\*.+?\*\*|`[^`]+`)")


def add_inline(p, text, size=12, ea=SONG, base_bold=False):
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("<sup>"):
            r = p.add_run(tok[5:-6].replace("\\*", "*"))
            set_run(r, size=size, bold=base_bold, ea=ea, sup=True)
        elif tok.startswith("**"):
            r = p.add_run(tok[2:-2].replace("\\*", "*"))
            set_run(r, size=size, bold=True, ea=ea)
        elif tok.startswith("`"):
            r = p.add_run(tok[1:-1])
            set_run(r, size=size, bold=base_bold, ea=ea)
        else:
            r = p.add_run(tok.replace("\\*", "*"))
            set_run(r, size=size, bold=base_bold, ea=ea)


def add_body(text, indent=True, align=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 2.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if indent:
        pf.first_line_indent = Pt(24)
    if align:
        p.alignment = align
    add_inline(p, text)
    return p


def add_heading_md(text, level):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    size = {1: 15, 2: 14, 3: 12}[level]
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, text, size=size, ea=HEI, base_bold=True)
    # outline level for navigation pane
    pPr = p._p.get_or_add_pPr()
    ol = pPr.makeelement(qn("w:outlineLvl"), {qn("w:val"): str(level - 1)})
    pPr.append(ol)
    return p


def add_table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    if not cells:
        return
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Table Grid"
    for i, row in enumerate(cells):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            add_inline(p, val, size=10.5, base_bold=(i == 0))
    doc.add_paragraph()


# ---------- render ----------
blocks = [b for b in re.split(r"\n\s*\n", text) if b.strip()]
current_section = ""
for b in blocks:
    raw_lines = b.splitlines()
    # markdown table block
    if all(l.lstrip().startswith("|") for l in raw_lines if l.strip()):
        add_table(raw_lines)
        continue
    for ln in raw_lines:
        s = ln.strip()
        if not s or s == "---":
            continue
        if s.startswith("### "):
            add_heading_md(s[4:], 3)
        elif s.startswith("## "):
            current_section = s[3:]
            add_heading_md(s[3:], 2)
        elif s.startswith("# "):
            current_section = s[2:]
            add_heading_md(s[2:], 1)
        elif s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 2.0
            add_inline(p, s[2:])
        elif current_section == "参考文献" and re.match(r"^\d+\.\s", s):
            add_body(s, indent=False)
        else:
            centered = s.startswith("**Development and benchmark")
            add_body(s, indent=not centered, align=WD_ALIGN_PARAGRAPH.CENTER if centered else None)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"[docx] {OUT} ({OUT.stat().st_size:,} bytes)")
