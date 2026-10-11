# -*- coding: utf-8 -*-
"""Build submission-ready EN Word file (manuscript.en.docx) from manuscript.en.md."""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "outputs" / "paper" / "manuscript.en.md"
OUT = ROOT / "outputs" / "submission" / "manuscript.en.docx"

ASCII, EA = "Times New Roman", "宋体"

text = SRC.read_text(encoding="utf-8")
lines = [ln for ln in text.splitlines() if not ln.strip().startswith(">")]
text = "\n".join(lines)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin = sec.bottom_margin = Mm(25.4)
sec.left_margin = sec.right_margin = Mm(31.7)
lnnum = sec._sectPr.makeelement(qn("w:lnNumType"), {qn("w:countBy"): "1", qn("w:restart"): "newPage"})
sec._sectPr.append(lnnum)

style = doc.styles["Normal"]
style.font.name = ASCII
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), EA)


def set_run(run, size=12, bold=False, sup=False):
    run.font.name = ASCII
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.superscript = sup
    rPr = run._element.get_or_add_rPr()
    rPr.get_or_add_rFonts().set(qn("w:eastAsia"), EA)


TOKEN = re.compile(r"(<sup>.+?</sup>|\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")


def add_inline(p, text, size=12, base_bold=False):
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("<sup>"):
            r = p.add_run(tok[5:-6].replace("\\*", "*"))
            set_run(r, size=size, bold=base_bold, sup=True)
        elif tok.startswith("**"):
            r = p.add_run(tok[2:-2].replace("\\*", "*"))
            set_run(r, size=size, bold=True)
        elif tok.startswith("`"):
            r = p.add_run(tok[1:-1])
            set_run(r, size=size, bold=base_bold)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = p.add_run(tok[1:-1].replace("\\*", "*"))
            r.font.italic = True
            set_run(r, size=size, bold=base_bold)
            r.font.italic = True
        else:
            r = p.add_run(tok.replace("\\*", "*"))
            set_run(r, size=size, bold=base_bold)


def add_body(text, indent=True, align=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 2.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if indent:
        pf.first_line_indent = Pt(36)
    if align:
        p.alignment = align
    add_inline(p, text)
    return p


def add_heading_md(text, level):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    size = {1: 15, 2: 14, 3: 12}[level]
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, text, size=size, base_bold=True)
    pPr = p._p.get_or_add_pPr()
    pPr.append(pPr.makeelement(qn("w:outlineLvl"), {qn("w:val"): str(level - 1)}))
    return p


def add_table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    if not cells:
        return
    t = doc.add_table(rows=len(cells), cols=len(cells[0]))
    t.style = "Table Grid"
    for i, row in enumerate(cells):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            add_inline(p, val, size=10.5, base_bold=(i == 0))
    doc.add_paragraph()


blocks = [b for b in re.split(r"\n\s*\n", text) if b.strip()]
current_section = ""
for b in blocks:
    raw_lines = b.splitlines()
    if all(l.lstrip().startswith("|") for l in raw_lines if l.strip()):
        add_table(raw_lines)
        continue
    for ln in raw_lines:
        s = ln.strip()
        if not s or s == "---":
            continue
        if s.startswith("### "):
            add_heading_md(s[4:], 3)
        elif s.startswith("## "):
            current_section = s[3:]
            add_heading_md(s[3:], 2)
        elif s.startswith("# "):
            current_section = s[2:]
            add_heading_md(s[2:], 1)
        elif s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.line_spacing = 2.0
            add_inline(p, s[2:])
        elif current_section.lower().startswith("references") and re.match(r"^\d+\.\s", s):
            add_body(s, indent=False)
        else:
            centered = s.startswith("# ") or s == ln.strip() and False
            add_body(s, indent=True)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(f"[docx] {OUT} ({OUT.stat().st_size:,} bytes)")
