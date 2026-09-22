# -*- coding: utf-8 -*-
"""Convert outputs/paper/manuscript.en.md to journal-submission Word format.

Format: A4, 2.5 cm margins, Times New Roman 12 pt, double spacing, true
Heading styles, figures embedded below their legends, hanging-indent
references. Run: .venv\\Scripts\\python.exe scripts\\export_manuscript_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "outputs" / "paper" / "manuscript.en.md"
OUT = ROOT / "outputs" / "submission" / "manuscript.en.docx"
FIG = ROOT / "outputs" / "figures"
LEGEND_FILES = {
    "Figure 1": "figure1_architecture.png",
    "Figure 2": "figure2_performance.png",
    "Figure 3": "figure3_expert_agreement.png",
    "Figure 4": "figure4_demo_case.png",
    "Figure 5": "figure5_reasoning_trace.png",
}

TNR = "Times New Roman"


def set_style(doc):
    n = doc.styles["Normal"]
    n.font.name = TNR
    n.font.size = Pt(12)
    pf = n.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_after = Pt(0)
    for name, size in [("Heading 1", 13), ("Heading 2", 12), ("Heading 3", 12)]:
        st = doc.styles[name]
        st.font.name = TNR
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(6)


TOKEN = re.compile(r"(<sup>.*?</sup>|<sub>.*?</sub>|\*\*.+?\*\*|\*[^*]+?\*)")


def add_runs(p, text):
    """Markdown-lite inline parser: sup/sub/bold/italic + unescape."""
    for piece in TOKEN.split(text):
        if not piece:
            continue
        sup = sub = bold = it = False
        if piece.startswith("<sup>"):
            piece, sup = piece[5:-6], True
        elif piece.startswith("<sub>"):
            piece, sub = piece[5:-6], True
        elif piece.startswith("**"):
            piece, bold = piece[2:-2], True
        elif piece.startswith("*"):
            piece, it = piece[1:-1], True
        piece = piece.replace("\\*", "*").replace("\\_", "_")
        r = p.add_run(piece)
        r.font.name = TNR
        r.font.superscript = sup
        r.font.subscript = sub
        r.bold = bold
        r.italic = it
    return p


def para(doc, text, *, align=None, size=None, bold=False, hanging=False):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    if hanging:
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
    add_runs(p, text)
    for r in p.runs:
        if size:
            r.font.size = Pt(size)
        if bold:
            r.bold = True
    return p


def main():
    doc = Document()
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Cm(2.5)
    set_style(doc)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i, in_refs = 0, False
    while i < len(lines):
        ln = lines[i].rstrip()
        i += 1
        if not ln.strip():
            continue
        if ln.startswith("# ") and not ln.startswith("## "):
            para(doc, ln[2:], align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True)
        elif ln.startswith("### "):
            doc.add_heading("", level=2).add_run("").bold = True
            h = doc.paragraphs[-1]
            add_runs(h, ln[4:])
        elif ln.startswith("## "):
            title = ln[3:]
            in_refs = title == "References"
            h = doc.add_heading("", level=1)
            add_runs(h, title)
        elif re.match(r"^\d+\.\s", ln) and in_refs:
            para(doc, ln, hanging=True)
        elif ln.startswith("- "):
            para(doc, ln[2:], hanging=True)
        else:
            p = para(doc, ln)
            # embed figure below its legend
            m = re.match(r"\*\*(Figure \d)\.", ln)
            if m and m.group(1) in LEGEND_FILES:
                img = FIG / LEGEND_FILES[m.group(1)]
                if img.exists():
                    fp = doc.add_paragraph()
                    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    fp.add_run().add_picture(str(img), width=Cm(16))
    doc.save(OUT)
    print(f"[docx] {OUT} ({OUT.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
